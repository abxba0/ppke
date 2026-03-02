"""Universal document converter — detects format and converts to Markdown.

All converters follow the same contract: accept a ``pathlib.Path`` and return a
Markdown string that the existing ``ppke.parser.markdown`` module can ingest.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Callable

logger = logging.getLogger(__name__)

# Mapping of lowercase file extension → converter function
_CONVERTERS: dict[str, Callable[[Path], str]] = {}


def register(ext: str):
    """Decorator to register a converter for a file extension."""

    def decorator(fn: Callable[[Path], str]):
        _CONVERTERS[ext.lower()] = fn
        return fn

    return decorator


# ------------------------------------------------------------------
# Built-in converters
# ------------------------------------------------------------------


@register(".md")
def _convert_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8")


@register(".txt")
def _convert_text(path: Path) -> str:
    """Wrap plain text in minimal Markdown heading structure."""
    text = path.read_text(encoding="utf-8")
    return f"# {path.stem}\n\n{text}"


@register(".pdf")
def _convert_pdf(path: Path) -> str:
    """Extract text from PDF using PyMuPDF (fitz).

    Falls back to OCR via pytesseract + pdf2image when pages yield no text.
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise ImportError(
            "PDF support requires PyMuPDF. Install with: pip install 'ppke[ocr]'"
        )

    doc = fitz.open(str(path))
    pages: list[str] = []
    ocr_pages: list[int] = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()
        if text:
            pages.append(f"## Page {page_num + 1}\n\n{text}")
        else:
            ocr_pages.append(page_num)
    doc.close()

    # OCR fallback for pages with no extractable text
    if ocr_pages:
        try:
            from ppke.converter.ocr import ocr_pdf_pages

            ocr_results = ocr_pdf_pages(path, ocr_pages)
            for page_num, text in ocr_results.items():
                pages.append(f"## Page {page_num + 1} (OCR)\n\n{text}")
        except ImportError:
            logger.warning(
                "OCR dependencies not installed — %d scanned pages skipped. "
                "Install with: pip install 'ppke[ocr]'",
                len(ocr_pages),
            )

    title = path.stem.replace("_", " ").replace("-", " ").title()
    return f"# {title}\n\n" + "\n\n".join(pages)


@register(".docx")
def _convert_docx(path: Path) -> str:
    """Convert Word document to Markdown preserving heading hierarchy."""
    try:
        from docx import Document  # python-docx
    except ImportError:
        raise ImportError(
            "DOCX support requires python-docx. Install with: pip install 'ppke[ocr]'"
        )

    doc = Document(str(path))
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower()
        if "heading 1" in style:
            lines.append(f"# {text}")
        elif "heading 2" in style:
            lines.append(f"## {text}")
        elif "heading 3" in style:
            lines.append(f"### {text}")
        elif "title" in style:
            lines.append(f"# {text}")
        else:
            lines.append(text)
        lines.append("")  # blank line between paragraphs

    return "\n".join(lines)


@register(".epub")
def _convert_epub(path: Path) -> str:
    """Convert EPUB ebook to Markdown."""
    try:
        import ebooklib  # noqa: F401
        from ebooklib import epub
    except ImportError:
        raise ImportError(
            "EPUB support requires ebooklib + beautifulsoup4. "
            "Install with: pip install ebooklib beautifulsoup4"
        )

    from bs4 import BeautifulSoup

    book = epub.read_epub(str(path))
    sections: list[str] = []

    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        text = soup.get_text(separator="\n").strip()
        if text:
            sections.append(text)

    title = path.stem.replace("_", " ").replace("-", " ").title()
    return f"# {title}\n\n" + "\n\n---\n\n".join(sections)


@register(".html")
@register(".htm")
def _convert_html(path: Path) -> str:
    """Convert HTML file to clean Markdown text."""
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        raise ImportError(
            "HTML support requires beautifulsoup4. "
            "Install with: pip install beautifulsoup4"
        )

    html = path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html, "html.parser")

    # Remove scripts and styles
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    text = soup.get_text(separator="\n").strip()
    title = soup.title.string if soup.title else path.stem
    return f"# {title}\n\n{text}"


# Image formats — OCR
for _ext in (".jpg", ".jpeg", ".png", ".tiff", ".tif", ".bmp", ".webp"):

    @register(_ext)
    def _convert_image(path: Path) -> str:
        """OCR an image file to Markdown."""
        from ppke.converter.ocr import ocr_image

        text = ocr_image(path)
        return f"# {path.stem}\n\n{text}"


@register(".pptx")
def _convert_pptx(path: Path) -> str:
    """Convert PowerPoint to Markdown (one heading per slide)."""
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError(
            "PPTX support requires python-pptx. Install with: pip install python-pptx"
        )

    prs = Presentation(str(path))
    slides: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = paragraph.text.strip()
                    if t:
                        texts.append(t)
        if texts:
            title = texts[0]
            body = "\n\n".join(texts[1:]) if len(texts) > 1 else ""
            slide_md = f"## Slide {i}: {title}"
            if body:
                slide_md += f"\n\n{body}"
            slides.append(slide_md)

    doc_title = path.stem.replace("_", " ").replace("-", " ").title()
    return f"# {doc_title}\n\n" + "\n\n".join(slides)


# Audio formats — transcription
for _ext in (".mp3", ".mp4", ".m4a", ".wav", ".webm", ".mpeg", ".mpga", ".ogg"):

    @register(_ext)
    def _convert_audio(path: Path) -> str:
        """Transcribe audio file to Markdown."""
        from ppke.audio.transcriber import transcribe

        return transcribe(path)


@register(".tex")
def _convert_latex(path: Path) -> str:
    """Convert LaTeX source to Markdown via pandoc (falls back to regex strip)."""
    import re
    import subprocess

    try:
        result = subprocess.run(
            ["pandoc", str(path), "-f", "latex", "-t", "markdown", "--wrap=none"],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode == 0 and result.stdout.strip():
            return f"# {path.stem}\n\n{result.stdout}"
    except (FileNotFoundError, subprocess.TimeoutExpired):
        logger.warning("pandoc not available or timed out — falling back to regex LaTeX strip")

    # Regex fallback: strip common LaTeX commands and recover text
    text = path.read_text(encoding="utf-8", errors="replace")
    # Remove preamble
    body_match = re.search(r"\\begin\{document\}(.*?)\\end\{document\}", text, re.DOTALL)
    if body_match:
        text = body_match.group(1)
    # \section{title} → ## title
    text = re.sub(r"\\(?:sub)*section\*?\{([^}]*)\}", lambda m: "## " + m.group(1), text)
    # \emph{x}, \textbf{x}, \textit{x} → x
    text = re.sub(r"\\(?:emph|textbf|textit|text)\{([^}]*)\}", r"\1", text)
    # Generic \cmd{arg} → arg
    text = re.sub(r"\\[a-zA-Z]+\{([^}]*)\}", r"\1", text)
    # Lone commands
    text = re.sub(r"\\[a-zA-Z]+\*?", "", text)
    # Remaining braces
    text = re.sub(r"[{}]", "", text)
    # Collapse whitespace
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return f"# {path.stem}\n\n{text}"


@register(".csv")
def _convert_csv(path: Path) -> str:
    """Convert a CSV file to a Markdown table (max 500 data rows)."""
    import csv

    rows: list[list[str]] = []
    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append(row)

    if not rows:
        return f"# {path.stem}\n\n*(empty file)*"

    header = rows[0]
    data_rows = rows[1:]
    col_count = len(header)

    lines: list[str] = [f"# {path.stem}", ""]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for row in data_rows[:500]:
        # Pad or trim to match header width
        padded = (row + [""] * col_count)[:col_count]
        lines.append("| " + " | ".join(c.replace("|", "\\|") for c in padded) + " |")
    if len(data_rows) > 500:
        lines.append(f"\n*({len(data_rows) - 500} additional rows not shown)*")

    return "\n".join(lines)


@register(".xlsx")
@register(".xls")
def _convert_excel(path: Path) -> str:
    """Convert an Excel workbook to Markdown tables (one section per sheet)."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "Excel support requires openpyxl. Install with: pip install openpyxl"
        )

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not open Excel file '{path.name}': {exc}") from exc

    sheets_md: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append([str(v) if v is not None else "" for v in row])

        if not rows:
            continue

        header = rows[0]
        data_rows = rows[1:]
        col_count = len(header)

        lines: list[str] = [f"## {sheet_name}", ""]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join("---" for _ in header) + " |")
        for row in data_rows[:500]:
            padded = (row + [""] * col_count)[:col_count]
            lines.append("| " + " | ".join(c.replace("|", "\\|") for c in padded) + " |")
        if len(data_rows) > 500:
            lines.append(f"\n*({len(data_rows) - 500} additional rows not shown)*")

        sheets_md.append("\n".join(lines))

    wb.close()

    title = path.stem.replace("_", " ").replace("-", " ").title()
    if not sheets_md:
        return f"# {title}\n\n*(empty workbook)*"
    return f"# {title}\n\n" + "\n\n".join(sheets_md)


@register(".bib")
def _convert_bibtex(path: Path) -> str:
    """Convert BibTeX bibliography file to Markdown."""
    from ppke.converter.zotero import convert_zotero_file

    return convert_zotero_file(path)


@register(".rdf")
def _convert_rdf(path: Path) -> str:
    """Convert Zotero RDF export to Markdown."""
    from ppke.converter.zotero import convert_zotero_file

    return convert_zotero_file(path)


@register(".zip")
def _convert_zip(path: Path) -> str:
    """Extract a ZIP archive and convert each supported file inside it.

    Returns a single combined Markdown document with each file as a section.
    Raises ``ValueError`` if the archive contains no supported files.
    """
    import zipfile

    results: list[str] = []

    with zipfile.ZipFile(str(path), "r") as zf:
        for member in zf.namelist():
            member_path = Path(member)
            ext = member_path.suffix.lower()

            # Skip directories, hidden files, and unsupported formats
            if member.endswith("/") or member_path.name.startswith("."):
                continue
            if ext not in _CONVERTERS:
                continue

            import tempfile

            with tempfile.TemporaryDirectory() as tmp_dir:
                extracted = Path(tmp_dir) / member_path.name
                extracted.write_bytes(zf.read(member))
                try:
                    content = _CONVERTERS[ext](extracted)
                    results.append(
                        f"---\n\n<!-- file: {member} -->\n\n{content}"
                    )
                except Exception as exc:
                    logger.warning("Skipping '%s' in ZIP: %s", member, exc)

    if not results:
        raise ValueError(
            f"ZIP archive '{path.name}' contains no supported files. "
            f"Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    title = path.stem.replace("_", " ").replace("-", " ").title()
    return f"# {title}\n\n" + "\n\n".join(results)


# ------------------------------------------------------------------
# Public API
# ------------------------------------------------------------------

SUPPORTED_EXTENSIONS: set[str] = set(_CONVERTERS.keys())


def convert_to_markdown(file_path: Path) -> str:
    """Convert any supported file to Markdown text.

    Raises ``ValueError`` for unsupported formats and ``ImportError`` if
    the required optional dependency is not installed.
    """
    ext = file_path.suffix.lower()
    converter = _CONVERTERS.get(ext)
    if converter is None:
        raise ValueError(
            f"Unsupported file format: '{ext}'. "
            f"Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    logger.info("Converting %s via %s converter", file_path.name, ext)
    return converter(file_path)
